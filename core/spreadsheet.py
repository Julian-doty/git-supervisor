"""
Writes the final repo analysis to .xlsx (openpyxl) or .ods (odfpy).

Deliberately dumb: this module never talks to GitHub or the model. It takes
rows that are already fully assembled (deterministic metadata + parsed
model judgement) and lays them out. Keeping file-writing separate from data
gathering means either backend script can call the same writer.
"""

from __future__ import annotations

from pathlib import Path

# Column header -> row dict key. Order here is the order columns appear in
# the output file.
COLUMNS: list[tuple[str, str]] = [
    ("Repo", "full_name"),
    ("URL", "url"),
    ("Stars", "stars"),
    ("Forks", "forks"),
    ("Language", "language"),
    ("License", "license_name"),
    ("Last Updated", "pushed_at"),
    ("Open Issues", "open_issues"),
    ("Topics", "topics"),
    ("Description", "description"),
    ("Relevance", "relevance"),
    ("Maintenance Read", "maintenance"),
    ("Implementation", "implementation"),
    ("Verdict", "verdict"),
]

# Columns that tend to hold long prose — given more width and word-wrap.
WIDE_COLUMNS = {"Description", "Relevance", "Maintenance Read", "Implementation"}


def _cell_value(row: dict, key: str):
    value = row.get(key, "")
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return value if value is not None else ""


def write_xlsx(rows: list[dict], out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Git Supervisor"

    for col_idx, (header, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (header, key) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=_cell_value(row, key))
            if header in WIDE_COLUMNS:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col_idx, (header, _) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = 50 if header in WIDE_COLUMNS else 18

    ws.freeze_panes = "A2"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def write_ods(rows: list[dict], out_path: Path) -> None:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableRow, TableCell, TableColumn
    from odf.text import P
    from odf.style import Style, TextProperties, TableColumnProperties

    doc = OpenDocumentSpreadsheet()

    bold_style = Style(name="HeaderBold", family="table-cell")
    bold_style.addElement(TextProperties(fontweight="bold"))
    doc.automaticstyles.addElement(bold_style)

    narrow_col = Style(name="ColNarrow", family="table-column")
    narrow_col.addElement(TableColumnProperties(columnwidth="2.2cm"))
    doc.automaticstyles.addElement(narrow_col)

    wide_col = Style(name="ColWide", family="table-column")
    wide_col.addElement(TableColumnProperties(columnwidth="7cm"))
    doc.automaticstyles.addElement(wide_col)

    table = Table(name="Git Supervisor")

    for header, _ in COLUMNS:
        col_style = wide_col if header in WIDE_COLUMNS else narrow_col
        table.addElement(TableColumn(stylename=col_style))

    header_row = TableRow()
    for header, _ in COLUMNS:
        cell = TableCell(stylename=bold_style)
        cell.addElement(P(text=header))
        header_row.addElement(cell)
    table.addElement(header_row)

    for row in rows:
        data_row = TableRow()
        for _, key in COLUMNS:
            cell = TableCell()
            value = _cell_value(row, key)
            cell.addElement(P(text=str(value)))
            data_row.addElement(cell)
        table.addElement(data_row)

    doc.spreadsheet.addElement(table)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def write_spreadsheet(rows: list[dict], out_path: Path, fmt: str) -> None:
    """fmt must be 'xlsx' or 'ods'."""
    if fmt == "xlsx":
        write_xlsx(rows, out_path)
    elif fmt == "ods":
        write_ods(rows, out_path)
    else:
        raise ValueError(f"Unsupported spreadsheet format: {fmt!r} (use 'xlsx' or 'ods')")
